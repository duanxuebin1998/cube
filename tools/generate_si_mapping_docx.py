from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/01_协议与寄存器/SI协议适配/02_协议映射/SI协议兼容映射表.md"
OUTPUT = ROOT / "docs/01_协议与寄存器/SI协议适配/02_协议映射/SI协议兼容映射表.docx"


TABLE_HEADER_FILL = "E8EEF5"
TABLE_BORDER_COLOR = "B8C2CC"
HEADING_BLUE = RGBColor(46, 116, 181)
HEADING_DARK = RGBColor(31, 77, 120)


def set_run_font(run, name: str = "Calibri", east_asia: str = "Microsoft YaHei", size: Pt | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = size
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), TABLE_BORDER_COLOR)


def set_cell_margins(table, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    values = {"top": top, "start": start, "bottom": bottom, "end": end}
    for key, value in values.items():
        element = margins.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    table.autofit = False


def split_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [cell.strip().replace("<br>", "\n") for cell in text.split("|")]


def is_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def parse_markdown(path: Path) -> list[dict]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            header = split_table_row(stripped)
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            blocks.append({"type": "heading", "level": len(heading.group(1)), "text": heading.group(2)})
            i += 1
            continue
        bullet = re.match(r"^(\s*)-\s+(.+)$", line)
        if bullet:
            indent = len(bullet.group(1)) // 2
            blocks.append({"type": "bullet", "level": indent, "text": bullet.group(2).strip()})
            i += 1
            continue
        numbered = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if numbered:
            indent = len(numbered.group(1)) // 2
            blocks.append({"type": "number", "level": indent, "text": numbered.group(2).strip()})
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("|")
                or re.match(r"^\s*-\s+", lines[i])
                or re.match(r"^\s*\d+\.\s+", lines[i])
            ):
                break
            paragraph_lines.append(nxt)
            i += 1
        blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines)})
    return blocks


def add_inline_text(paragraph, text: str, bold_default: bool = False) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", east_asia="Microsoft YaHei")
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run)
    for run in paragraph.runs:
        if bold_default:
            run.bold = True


def apply_document_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, HEADING_BLUE, 18, 10),
        ("Heading 2", 13, HEADING_BLUE, 14, 7),
        ("Heading 3", 12, HEADING_DARK, 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("SI协议兼容映射表")
    set_run_font(run, size=Pt(22))
    run.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("维护版 Word 文档；内容来源为当前 Markdown 映射表和官方 Modbus 手册核对口径。")
    set_run_font(run, size=Pt(10.5))
    run.font.color.rgb = RGBColor(85, 85, 85)


def choose_widths(headers: list[str], row_count: int) -> list[int]:
    n = len(headers)
    if n == 2:
        return [2500, 6860]
    if n == 3:
        return [1700, 3900, 3760]
    if n == 4:
        return [1500, 2900, 2500, 2460]
    if n == 5:
        return [1300, 1800, 2600, 1500, 2160]
    return [9360 // n for _ in range(n)]


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    if not header:
        return
    col_count = len(header)
    table = doc.add_table(rows=1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    set_table_borders(table)
    set_cell_margins(table)
    widths = choose_widths(header, len(rows))

    def format_cell(cell, text: str, width_dxa: int, is_header: bool = False) -> None:
        cell.width = Inches(width_dxa / 1440)
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(width_dxa))
        tc_w.set(qn("w:type"), "dxa")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if is_header:
            set_cell_shading(cell, TABLE_HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        add_inline_text(paragraph, text, bold_default=is_header)
        for run in paragraph.runs:
            run.font.size = Pt(9 if col_count >= 5 else 9.5)

    for idx, cell in enumerate(table.rows[0].cells):
        format_cell(cell, header[idx], widths[idx], True)
    for row in rows:
        cells = table.add_row().cells
        for idx in range(col_count):
            text = row[idx] if idx < len(row) else ""
            format_cell(cells[idx], text, widths[idx], False)
    doc.add_paragraph()


def build_docx() -> None:
    blocks = parse_markdown(SOURCE)
    doc = Document()
    apply_document_styles(doc)
    add_title_block(doc)

    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            text = block["text"]
            if level == 1:
                continue
            para = doc.add_heading(text, level=min(level - 1, 3))
            for run in para.runs:
                set_run_font(run, size=run.font.size)
        elif kind == "paragraph":
            para = doc.add_paragraph()
            add_inline_text(para, block["text"])
        elif kind == "bullet":
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.left_indent = Cm(0.7 + block["level"] * 0.55)
            para.paragraph_format.first_line_indent = Cm(-0.35)
            add_inline_text(para, block["text"])
        elif kind == "number":
            para = doc.add_paragraph(style="List Number")
            para.paragraph_format.left_indent = Cm(0.7 + block["level"] * 0.55)
            para.paragraph_format.first_line_indent = Cm(-0.35)
            add_inline_text(para, block["text"])
        elif kind == "table":
            add_table(doc, block["header"], block["rows"])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("SI协议兼容映射表")
    set_run_font(run, size=Pt(9))
    run.font.color.rgb = RGBColor(85, 85, 85)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT)
